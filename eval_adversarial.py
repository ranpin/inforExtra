"""对抗子集评测：测 laya 管线在「分布外」的召回天花板与否定泛化。

三族各 ~20 条，全部针对 jev_extract 的已知代码盲区设计：
  族1 表外姓   —— 姓氏不在 SURNAME_CHARS / 外号非「小X」，surname 过滤生成不出候选 → 预期漏召(FN)
  族2 真实否定 —— 否定线索不在 NEGATION_CUES(不是/不叫/别叫/别说/别喊/不姓) 或置于名字之后 → 预期误报(FP)
  族3 ASR噪声  —— 口语填充词/重复/语气尾词，人名仍在 → 预期基本保持（对照组）

v4 规则优化（2026-10-21）后：族1 F1 0.824（生僻姓/阿X/称谓后缀锚点修复，
余 6FN 为设计内排除）、族2 FP 15→1、族3 1.000，合计 F1 0.474→0.907。

gold 由构造时确定（构造即标注）。指标 = 人名级 P/R/F1（与主评测同口径）。
产物：test/对抗集.xlsx（输入）/ test/对抗集_gold.xlsx（gold）/ test/results/对抗集_laya_v4.xlsx（预测）
用法：.venv/bin/python eval_adversarial.py
"""
import os
os.environ.setdefault("JEV_BACKEND", "laya")
os.environ.setdefault("JEV_CANDIDATE_FILTER", "surname")
os.environ.setdefault("JEV_ASSEMBLY", "forced")
os.environ.setdefault("HF_HUB_OFFLINE", "1")
os.environ.setdefault("TRANSFORMERS_OFFLINE", "1")

import json
import jev_extract
import openpyxl

# ---- 族1：表外姓 / 外号（姓氏不在 SURNAME_CHARS，预期漏召）----
FAM1 = [
    ("副驾坐的是翦老师", [{"name": "翦老师", "location": "副驾"}]),
    ("主驾是仝工", [{"name": "仝工", "location": "主驾"}]),
    ("后排左边是冼医生", [{"name": "冼医生", "location": "后排左边"}]),
    ("我是佟老板，坐在主驾", [{"name": "佟老板", "location": "主驾"}]),
    ("副驾是佘经理", [{"name": "佘经理", "location": "副驾"}]),
    ("后排坐着菅总", [{"name": "菅总", "location": "后排"}]),
    ("主驾是庹师傅", [{"name": "庹师傅", "location": "主驾"}]),
    ("后排右边是商总", [{"name": "商总", "location": "后排右边"}]),
    ("主驾是言老师", [{"name": "言老师", "location": "主驾"}]),
    ("副驾坐着阿宝", [{"name": "阿宝", "location": "副驾"}]),
    ("后排是大壮", [{"name": "大壮", "location": "后排"}]),
    ("主驾是阿强", [{"name": "阿强", "location": "主驾"}]),
    ("副驾是胖子", [{"name": "胖子", "location": "副驾"}]),
    ("后排左边坐着初工", [{"name": "初工", "location": "后排左边"}]),
    ("主驾是展总", [{"name": "展总", "location": "主驾"}]),
    ("副驾是迟医生", [{"name": "迟医生", "location": "副驾"}]),
    ("后排是慕老师", [{"name": "慕老师", "location": "后排"}]),
    ("主驾坐着福老板", [{"name": "福老板", "location": "主驾"}]),
    ("副驾是贵经理", [{"name": "贵经理", "location": "副驾"}]),
    ("后排右边是撒师傅", [{"name": "撒师傅", "location": "后排右边"}]),
]

# ---- 族2：真实否定（线索不在 NEGATION_CUES 或在名字后，预期误报）----
FAM2 = [
    ("别把我当成张医生", []),
    ("陈总不是我，他还没上车", []),
    ("我并非李经理", []),
    ("我没说我是王老板", []),
    ("搞错了，副驾并非小刘", []),
    ("副驾那位算不上张总", []),
    ("我哪里是李医生啊", []),
    ("别误会，我并非王师傅", []),
    ("刚才提到的赵总并不在车上", []),
    ("周老师没来", []),
    ("吴经理今天缺席", []),
    ("郑总临时不来了", []),
    ("我说的是别人，并非孙工", []),
    ("别记成马老板", []),
    ("胡老师是别人，不是我", []),
    ("别搞混了，林总不在车上", []),
    ("何医生今天休息，没来", []),
    ("朱这个姓跟我没关系", []),
    ("高总临时有事不来了", []),
    ("罗老师并不是我", []),
]

# ---- 族3：ASR噪声/口语（人名仍在，预期基本保持，对照组）----
FAM3 = [
    ("那个...副驾是小王吧", [{"name": "小王", "location": "副驾"}]),
    ("呃后排坐的是陈总", [{"name": "陈总", "location": "后排"}]),
    ("副驾是...是小刘", [{"name": "小刘", "location": "副驾"}]),
    ("主驾是王总王总", [{"name": "王总", "location": "主驾"}]),
    ("后排左边是张医生对吧", [{"name": "张医生", "location": "后排左边"}]),
    ("嗯主驾坐着李经理", [{"name": "李经理", "location": "主驾"}]),
    ("副驾嘛就是赵老板", [{"name": "赵老板", "location": "副驾"}]),
    ("那个谁，后排是孙工", [{"name": "孙工", "location": "后排"}]),
    ("主驾是周总周总没错", [{"name": "周总", "location": "主驾"}]),
    ("副驾是小吴哈", [{"name": "小吴", "location": "副驾"}]),
    ("后排右边坐着郑经理呢", [{"name": "郑经理", "location": "后排右边"}]),
    ("诶主驾是冯老师", [{"name": "冯老师", "location": "主驾"}]),
    ("副驾是陈总，对，陈总", [{"name": "陈总", "location": "副驾"}]),
    ("后排是蒋工吧好像", [{"name": "蒋工", "location": "后排"}]),
    ("主驾坐着沈老板嗯", [{"name": "沈老板", "location": "主驾"}]),
    ("副驾是韩医生没错吧", [{"name": "韩医生", "location": "副驾"}]),
    ("那个后排左边是杨总", [{"name": "杨总", "location": "后排左边"}]),
    ("主驾是朱师傅呀", [{"name": "朱师傅", "location": "主驾"}]),
    ("副驾...副驾是秦经理", [{"name": "秦经理", "location": "副驾"}]),
    ("后排坐着许老板哈", [{"name": "许老板", "location": "后排"}]),
]

FAMILIES = [("表外姓", FAM1), ("真实否定", FAM2), ("ASR噪声", FAM3)]


def prf(tp, fp, fn):
    p = tp / (tp + fp) if tp + fp else 0.0
    r = tp / (tp + fn) if tp + fn else 0.0
    f = 2 * p * r / (p + r) if p + r else 0.0
    return p, r, f


def main():
    all_rows, summary = [], []
    gtp = gfp = gfn = 0
    for fam, examples in FAMILIES:
        tp = fp = fn = 0
        for text, gold in examples:
            r = jev_extract.extract_entities(text)
            pred = r.get("persons", [])
            pn = {p["name"] for p in pred}
            gn = {g["name"] for g in gold}
            tp += len(pn & gn); fp += len(pn - gn); fn += len(gn - pn)
            all_rows.append((fam, text, json.dumps(gold, ensure_ascii=False),
                             json.dumps(pred, ensure_ascii=False)))
        p, rr, f = prf(tp, fp, fn)
        summary.append((fam, len(examples), p, rr, f, tp, fp, fn))
        gtp += tp; gfp += fp; gfn += fn

    gp, gr, gf = prf(gtp, gfp, gfn)
    print("\n" + "=" * 62)
    print("对抗子集评测（laya 最优配置，人名级 P/R/F1）")
    print("=" * 62)
    print(f"{'族':<10}{'条数':>5}{'P':>8}{'R':>8}{'F1':>8}{'TP':>5}{'FP':>5}{'FN':>5}")
    for fam, n, p, rr, f, tp, fp, fn in summary:
        print(f"{fam:<10}{n:>5}{p:>8.3f}{rr:>8.3f}{f:>8.3f}{tp:>5}{fp:>5}{fn:>5}")
    print("-" * 62)
    print(f"{'合计':<10}{sum(s[1] for s in summary):>5}{gp:>8.3f}{gr:>8.3f}{gf:>8.3f}{gtp:>5}{gfp:>5}{gfn:>5}")
    print("=" * 62)
    print("对照：分布内 510 条 v4 人名 F1 = 0.944")

    # 产物入库
    os.makedirs("test/results", exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "对抗集"
    ws.append(["用户输入"]); 
    for _, ex in FAMILIES:
        for t, _ in ex: ws.append([t])
    wb.save("test/对抗集.xlsx")

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "gold"
    ws.append(["用户输入", "JSON输出"])
    for _, ex in FAMILIES:
        for t, g in ex: ws.append([t, json.dumps(g, ensure_ascii=False)])
    wb.save("test/对抗集_gold.xlsx")

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "预测"
    ws.append(["族", "用户输入", "gold", "预测"])
    for row in all_rows: ws.append(list(row))
    wb.save("test/results/对抗集_laya_v4.xlsx")
    print("\n产物已写入 test/对抗集.xlsx, test/对抗集_gold.xlsx, test/results/对抗集_laya_v4.xlsx")


if __name__ == "__main__":
    main()
